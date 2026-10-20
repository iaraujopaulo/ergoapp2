"""Exportação do ErgoApp: PDF estruturado (reportlab) e Excel dinâmico (pandas + openpyxl)."""
import io
from datetime import datetime
import cv2, pandas as pd
from openpyxl.chart import LineChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

RISCO = {1: "Aceitável", 2: "Aceitável", 3: "Investigar mais", 4: "Investigar mais",
         5: "Investigar e mudar em breve", 6: "Investigar e mudar em breve", 7: "Investigar e mudar já"}
CONCLUSAO = {1: "Postura aceitável, desde que não seja mantida ou repetida por longos períodos.",
             3: "São necessárias investigações adicionais; mudanças podem ser necessárias.",
             5: "São necessárias investigação e mudanças em breve.",
             7: "São necessárias investigação e mudanças imediatas."}
for _k, _b in ((2, 1), (4, 3), (6, 5)): CONCLUSAO[_k] = CONCLUSAO[_b]
COR = {1: "#2e9e5b", 2: "#2e9e5b", 3: "#d9a800", 4: "#d9a800", 5: "#e07b00", 6: "#e07b00", 7: "#d62839"}
COLS = ["tempo_s", "vista", "rula", "escore_A", "escore_B", "braco", "antebraco", "punho", "pescoco", "tronco", "pernas"]
ANG = {"braco": "Flexão do braço", "abducao": "Abdução do braço", "cotovelo": "Cotovelo",
       "punho": "Desvio do punho", "pescoco": "Flexão do pescoço", "tronco": "Flexão do tronco"}

def linha_hist(t, r):
    return dict(zip(COLS, [round(t, 2), r['view'], r['final'], r['scoreA'], r['scoreB'],
                           r['ua'], r['la'], r['wr'], r['neck'], r['trunk'], r['legs']]))

def passos(r):
    return [("Braço (com ajustes)", r['ua']), ("Antebraço", r['la']), ("Punho (com desvio)", r['wr']),
            ("Giro do punho", r['tw']), ("Tabela A", r['A']), ("Uso muscular", r['muscle']),
            ("Força/carga", r['force']), ("Escore A (Tabela A + uso + carga)", r['scoreA']),
            ("Pescoço", r['neck']), ("Tronco", r['trunk']), ("Pernas", r['legs']), ("Tabela B", r['B']),
            ("Escore B (Tabela B + uso + carga)", r['scoreB']), ("Escore final - Tabela C (A x B)", r['final'])]

def estat(h):
    if h is None or len(h) < 2: return None
    return dict(n=len(h), max=int(h.rula.max()), media=round(float(h.rula.mean()), 2), pct=round(100*float((h.rula >= 5).mean())))

def recomendacoes(r):
    l = []
    if r['ua'] >= 3: l.append("Braço: reduzir elevação e abdução, aproximar a carga do corpo, ajustar a altura da bancada e usar apoio de antebraço.")
    if r['la'] >= 2: l.append("Antebraço: manter o cotovelo entre 60 e 100 graus e evitar cruzar a linha média do corpo.")
    if r['wr'] >= 3: l.append("Punho: manter posição neutra e revisar ferramentas, teclado e mouse.")
    if r['neck'] >= 3: l.append("Pescoço: elevar o plano de visão (monitor/peça) e evitar flexão e giro prolongados.")
    if r['trunk'] >= 3: l.append("Tronco: ajustar altura e alcance, usar assento com encosto e evitar inclinação e torção.")
    if r['legs'] >= 2: l.append("Pernas: garantir apoio estável dos dois pés e distribuição de peso.")
    if r['force'] > 0: l.append("Carga: reduzir peso/força, usar auxílio mecânico e dividir a tarefa.")
    if r['muscle']: l.append("Uso muscular: incluir pausas e rodízio para quebrar a postura estática ou repetitiva.")
    return l or ["Manter as boas práticas atuais e reavaliar sempre que houver mudança no posto ou na tarefa."]

def excel_bytes(r, ident, hist):
    hist = hist if hist is not None else pd.DataFrame(columns=COLS); n = len(hist); buf = io.BytesIO()
    resumo = pd.DataFrame([("Data", datetime.now().strftime("%d/%m/%Y %H:%M")), *ident.items(), ("Vista", r['view']),
        ("Lado avaliado", r['side'].upper()), ("Escore final RULA", r['final']), ("Classificação", RISCO[r['final']]),
        ("Conclusão", CONCLUSAO[r['final']])], columns=["Campo", "Valor"])
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        resumo.to_excel(xw, sheet_name="Resumo", index=False)
        pd.DataFrame(passos(r), columns=["Etapa", "Pontos"]).to_excel(xw, sheet_name="Cruzamento A-B-C", index=False)
        pd.DataFrame([(ANG[k], v) for k, v in r['angles'].items() if k in ANG], columns=["Medida", "Graus"]).to_excel(xw, sheet_name="Ângulos", index=False)
        hist.to_excel(xw, sheet_name="Histórico", index=False)
        from ergoapp import TC
        pd.DataFrame(TC, index=range(1, 9), columns=range(1, 8)).rename_axis("A / B").to_excel(xw, sheet_name="Tabela C")
        ws, wh = xw.sheets["Resumo"], xw.sheets["Histórico"]
        ws.column_dimensions["A"].width = 34; ws.column_dimensions["B"].width = 70
        if n:   # fórmulas vivas: recalculam se o histórico for editado
            b = ws.max_row + 2
            for k, (lab, f) in enumerate([("RULA máximo (fórmula)", "=MAX('Histórico'!C:C)"),
                                          ("RULA médio (fórmula)", "=ROUND(AVERAGE('Histórico'!C:C),2)"),
                                          ("% quadros com RULA 5 ou mais", "=COUNTIF('Histórico'!C:C,\">=5\")/COUNT('Histórico'!C:C)")]):
                ws.cell(b+k, 1, lab); ws.cell(b+k, 2, f)
            ws.cell(b+2, 2).number_format = "0%"
            for lim, cor in ((7, "FF9AA5"), (5, "FFC58A"), (3, "FFE69A"), (1, "B7E4C7")):
                wh.conditional_formatting.add(f"C2:C{n+1}", CellIsRule(operator="greaterThanOrEqual", formula=[str(lim)],
                    fill=PatternFill("solid", start_color=cor, end_color=cor), stopIfTrue=True))
            wh.auto_filter.ref = wh.dimensions
            ch = LineChart(); ch.title = "RULA ao longo do tempo"; ch.y_axis.title = "Escore"; ch.x_axis.title = "Tempo (s)"
            ch.y_axis.scaling.min, ch.y_axis.scaling.max = 0, 7; ch.width, ch.height = 24, 9
            ch.add_data(Reference(wh, min_col=3, min_row=1, max_row=n+1), titles_from_data=True)
            ch.set_categories(Reference(wh, min_col=1, min_row=2, max_row=n+1)); wh.add_chart(ch, "M2")
    return buf.getvalue()

def _rodape(c, d):
    c.saveState(); c.setFont("Helvetica", 8); c.setFillColor(colors.grey)
    c.drawString(2*cm, 1.1*cm, "ErgoApp - análise postural RULA"); c.drawRightString(A4[0]-2*cm, 1.1*cm, f"Página {d.page}"); c.restoreState()

def _tabela(rows, widths, head=None):
    t = Table(([list(head)] if head else []) + [[str(c) for c in x] for x in rows], colWidths=widths, repeatRows=1 if head else 0)
    st = [("FONTSIZE", (0, 0), (-1, -1), 9), ("GRID", (0, 0), (-1, -1), .4, colors.HexColor("#c9ccd9")),
          ("ROWBACKGROUNDS", (0, 1 if head else 0), (-1, -1), [colors.white, colors.HexColor("#f3f4fa")])]
    if head: st += [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#3b3f6b")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold")]
    t.setStyle(TableStyle(st)); return t

def pdf_bytes(r, ident, hist=None, frame=None, laudo=False):
    buf = io.BytesIO(); s = estat(hist); ss = getSampleStyleSheet(); H1, H2, B = ss["Title"], ss["Heading2"], ss["BodyText"]
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=2*cm, rightMargin=2*cm, topMargin=1.8*cm, bottomMargin=2*cm,
                            title="Laudo ergonômico RULA" if laudo else "Relatório RULA")
    num = iter(range(1, 9)); sec = lambda t: Paragraph(f"{next(num)}. {t}" if laudo else t, H2)
    el = [Paragraph("LAUDO ERGONÔMICO - MÉTODO RULA" if laudo else "RELATÓRIO DE ANÁLISE POSTURAL - RULA", H1), sec("Identificação"),
          _tabela([("Data", datetime.now().strftime("%d/%m/%Y %H:%M")), *ident.items()], [5*cm, 12*cm])]
    if laudo:
        el += [sec("Metodologia"), Paragraph("Avaliação postural pelo método RULA (Rapid Upper Limb Assessment - McAtamney e Corlett, 1993), "
            "com detecção de pontos anatômicos por visão computacional (MediaPipe Pose) e cálculo de ângulos em 3D. As posturas são pontuadas e "
            "cruzadas nas Tabelas A (braço, antebraço, punho e giro), B (pescoço, tronco e pernas) e C, resultando no escore final de 1 a 7.", B)]
    el.append(sec("Resultados"))
    if frame is not None:
        h, w = frame.shape[:2]; k = min(9*cm/w, 10*cm/h); el.append(Image(io.BytesIO(cv2.imencode(".jpg", frame)[1].tobytes()), width=w*k, height=h*k))
    el += [Spacer(1, 6), _tabela([("Vista / lado avaliado", f"{r['view']} / {r['side'].upper()}"), *passos(r)], [11*cm, 6*cm], ("Componente", "Pontos"))]
    if s: el.append(Paragraph(f"Análise de {s['n']} quadros: RULA máximo {s['max']}, médio {s['media']}; {s['pct']}% dos quadros com escore igual ou superior a 5.", B))
    el += [sec("Classificação e conclusão"), Paragraph(f'<font color="{COR[r["final"]]}"><b>Nível {r["final"]} - {RISCO[r["final"]]}</b></font>', B),
           Paragraph(CONCLUSAO[r['final']], B)]
    if laudo:
        el.append(sec("Recomendações")); el += [Paragraph("- " + x, B) for x in recomendacoes(r)]
        el += [sec("Limitações"), Paragraph("Análise automatizada a partir de imagem 2D: o giro do punho e o desvio radial/ulnar não são medidos pelo modelo; "
            "ombro elevado, braço apoiado, carga e uso muscular são informados manualmente. O resultado é uma triagem e deve ser validado por "
            "profissional habilitado, podendo ser complementado por Análise Ergonômica do Trabalho (NR-17).", B), Spacer(1, 40),
            Paragraph("______________________________________________", B), Paragraph("Responsável técnico (nome, formação e registro profissional)", B)]
    doc.build(el, onFirstPage=_rodape, onLaterPages=_rodape); return buf.getvalue()
